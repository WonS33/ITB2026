"""제조내역 엑셀 업로드 규칙 (specs/05-production.md 4절)

샘플 엑셀(시트 "2026(NTI)") 구조
- 위쪽 몇 줄은 작성방법 안내, 그 아래 제목이 두 줄(병합)이다 → "일련번호" 제목이 있는 줄을 찾아 그 아래부터 읽는다
- 셀에는 짧은 값만 있고 앞부분은 셀 표시 형식에 있다
    일련번호 '0001' + 형식 "N8226"@  → N82260001
    모델명   '04'   + 형식 "NTI-"@   → NTI-04
    검사번호 '001'  + 형식 "NTK26"@  → NTK26001
- 일련번호가 없는 행(오른쪽 OPUN 등 숫자만 있는 행)은 제외한다
"""

import re
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from production.models import ManufacturingRecord

from .base import UploadHandler, parse_iso_date

# 엑셀 제목(공백 제거) → 저장할 필드
HEADER_MAP = {
    "일련번호": "serial",
    "모델명": "model_name",
    "부시가공": "bush_machining",
    "rpm": "rpm",
    "전원": "power",
    "제조년월": "manufactured_date",
    "검사일자": "inspection_date",
    "검사번호": "inspection_no",
    "PO번호": "po_no",
    "SO번호": "so_no",
    "거래처주문번호": "customer_order_no",
    "거래처": "customer",
    "최종발주처": "final_orderer",
    "사업명": "project_name",
    "납기일": "due_date",
    "출고일": "shipped_date",
    "비고": "note",
}
DATE_FIELDS = ("manufactured_date", "inspection_date", "due_date", "shipped_date")
SERIAL_PATTERN = re.compile(r"^(NE|N|M|B|Q|C|P)\d{8}$")  # 타입 문자 + 국가 2 + 년도 2 + 번호 4
TEXT_PREFIX = re.compile(r'^"([^"]*)"@$')  # 예: "N8226"@
NUMBER_PREFIX = re.compile(r'^"([^"]*)"(0+)$')  # 예: "NTK26"000


def _norm(value):
    return re.sub(r"\s+", "", str(value)) if value is not None else ""


def _number_text(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value)


def _format_prefix(cell):
    """셀 표시 형식 안의 앞 글자. 예: "N8226"@ → N8226"""
    fmt = cell.number_format or ""
    m = TEXT_PREFIX.match(fmt) or NUMBER_PREFIX.match(fmt)
    return m.group(1) if m else ""


def _display_text(cell):
    """엑셀 화면에 보이는 글자에 가깝게 바꾼다 (표시 형식의 앞 글자 + 셀 값)"""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    fmt = cell.number_format or ""
    m = NUMBER_PREFIX.match(fmt)
    if m and isinstance(value, (int, float)):
        return m.group(1) + _number_text(value).zfill(len(m.group(2)))
    text = _number_text(value).strip() if isinstance(value, (int, float)) else str(value).strip()
    m = TEXT_PREFIX.match(fmt)
    return m.group(1) + text if m else text


def _date_text(cell):
    value = cell.value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # 날짜가 숫자(예: 46031)로 저장된 경우 → 2026-01-09
        return from_excel(value).date().isoformat()
    return _display_text(cell)


def _serial(cell):
    """(전체 일련번호, 4자리) — 셀 값은 4자리, 앞부분은 표시 형식에서 가져온다"""
    value = cell.value
    short = (_number_text(value).zfill(4) if isinstance(value, (int, float)) else str(value).strip())
    prefix = _format_prefix(cell)
    full = prefix + short if prefix else short
    return full, full[-4:]


class ManufacturingExcelHandler(UploadHandler):
    kind = "manufacturing_excel"
    team = "production"
    allowed_extensions = (".xlsx",)
    key_field = "serial_full"
    columns = [
        ("serial_full", "일련번호"),
        ("model_name", "모델명"),
        ("bush_machining", "부시 가공"),
        ("rpm", "rpm"),
        ("power", "전원"),
        ("manufactured_date", "제조년월"),
        ("inspection_date", "검사일자"),
        ("inspection_no", "검사번호"),
        ("po_no", "PO번호"),
        ("so_no", "SO번호"),
        ("customer_order_no", "거래처 주문번호"),
        ("customer", "거래처"),
        ("final_orderer", "최종발주처"),
        ("project_name", "사업명"),
        ("due_date", "납기일"),
        ("shipped_date", "출고일"),
        ("note", "비고"),
    ]

    def parse(self, file):
        # data_only=True: 수식 셀은 계산된 값으로 읽는다
        workbook = load_workbook(file, data_only=True)
        rows, skipped = [], 0
        found_header = False
        for sheet in workbook.worksheets:
            result = self._parse_sheet(sheet)
            if result is None:
                continue
            found_header = True
            rows += result[0]
            skipped += result[1]
        if not found_header:
            raise ValueError("'일련번호' 제목 칸을 찾을 수 없습니다. 제조내역 엑셀 양식인지 확인해 주세요.")
        return rows, skipped

    def _parse_sheet(self, sheet):
        header_row = None
        for r in range(1, min(sheet.max_row, 30) + 1):
            if any(_norm(sheet.cell(r, c).value) == "일련번호" for c in range(1, min(sheet.max_column, 40) + 1)):
                header_row = r
                break
        if header_row is None:
            return None

        # 제목이 두 줄이므로 아래 줄(예: PO번호, SO번호)을 먼저 보고, 없으면 위 줄 제목을 쓴다.
        # 같은 제목이 두 번 나오면(예: 비고) 왼쪽 열을 쓴다.
        col_of, labeled_cols, sub_header_used = {}, set(), False
        for c in range(1, sheet.max_column + 1):
            upper = _norm(sheet.cell(header_row, c).value)
            lower = _norm(sheet.cell(header_row + 1, c).value)
            if upper or lower:
                labeled_cols.add(c)
            for name in (lower, upper):
                field = HEADER_MAP.get(name) or HEADER_MAP.get(name.lower() if name else "")
                if field and field not in col_of:
                    col_of[field] = c
                    sub_header_used |= name == lower
                    break

        rows, skipped = [], 0
        start = header_row + (2 if sub_header_used else 1)
        for r in range(start, sheet.max_row + 1):
            serial_cell = sheet.cell(r, col_of["serial"])
            if serial_cell.value in (None, ""):
                # 제목이 있는 열에 값이 하나라도 있으면 "일련번호 없는 행"으로 세고 제외한다
                if any(sheet.cell(r, c).value not in (None, "") for c in labeled_cols):
                    skipped += 1
                continue

            full, _ = _serial(serial_cell)
            values = {"serial_full": full}
            for field, c in col_of.items():
                if field == "serial":
                    continue
                cell = sheet.cell(r, c)
                if field in DATE_FIELDS:
                    values[field] = _date_text(cell)
                elif field == "bush_machining" and isinstance(cell.value, (int, float)):
                    values[field] = _number_text(cell.value).zfill(6)  # 301008 (표시: Ø30x10x08)
                elif field == "po_no":
                    values[field] = _display_text(cell).lstrip("#")  # #26006 → 26006
                else:
                    values[field] = _display_text(cell)
            for key, *_ in self.columns:
                values.setdefault(key, "")
            rows.append({"row_no": r, "sheet": sheet.title, "values": values})
        return rows, skipped

    def validate(self, values):
        errors = {}
        serial = values.get("serial_full", "").strip()
        if not serial:
            errors["serial_full"] = "일련번호가 없습니다."
        elif not SERIAL_PATTERN.match(serial):
            errors["serial_full"] = "형식: 타입 문자(N/M/B/Q/C/P/NE) + 숫자 8자리 (예: N82260001)"

        for field in DATE_FIELDS:
            try:
                if parse_iso_date(values.get(field)) is None and field == "manufactured_date":
                    errors[field] = "제조년월이 없습니다 (생산량 집계 기준)."
            except ValueError:
                errors[field] = "날짜 형식: YYYY-MM-DD"

        bush = values.get("bush_machining", "").strip()
        if bush and not re.fullmatch(r"\d{6}", bush):
            errors["bush_machining"] = "숫자 6자리 (예: 301008)"

        for key, *_ in self.columns:
            limit = ManufacturingRecord._meta.get_field(key).max_length
            if limit and len(values.get(key, "")) > limit and key not in errors:
                errors[key] = f"{limit}자 이하로 입력해 주세요."
        return errors

    def existing_keys(self, keys):
        return set(ManufacturingRecord.objects.filter(serial_full__in=keys).values_list("serial_full", flat=True))

    def save(self, rows, user):
        records = []
        for row in rows:
            v = {k: (val or "").strip() for k, val in row["values"].items()}
            for field in DATE_FIELDS:
                v[field] = parse_iso_date(v[field])
            records.append(
                ManufacturingRecord(
                    **{k: v[k] for k, *_ in self.columns},
                    serial_short=v["serial_full"][-4:],
                    source_sheet=row.get("sheet", ""),
                    created_by=user,
                )
            )
        ManufacturingRecord.objects.bulk_create(records)
        return len(records)
