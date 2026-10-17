"""가상 샘플 파일 만들기 (테스트·시연용, 실제 기업 데이터 아님)

실제 제조내역 엑셀과 같은 구조로 만든다: 안내 줄, 두 줄 제목(병합), 표시 형식으로 붙는 앞 글자,
일련번호 없는 행(OPUN 등 숫자만 있는 행).
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook

# 엑셀 열 위치 (샘플과 같은 순서)
HEADERS = ["일련번호", "모델명", "부시 가공", "rpm", "전원", "제조년월", "검사일자", "검사번호", "뉴토크 번호", None,
           "거래처 주문번호", "거래처", "최종발주처", "사업명", "납기일", "출고일", "비고"]


def build_manufacturing_workbook(rows, year=2026, extra_blank_rows=3, date_as_number=False):
    """rows: [{"no": 1, "model": "04", "made": date(...), ...}, ...] → xlsx 바이트"""
    yy = str(year)[2:]
    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}(NTI)"
    ws["B1"] = "** 작성방법(필독!!)"
    ws["D1"] = "일련번호: "
    ws["E1"] = "타입 문자 + 국가번호 2자리 + 년도 2자리 + 일련번호 4자리"
    ws["B7"] = f"{year}년 "

    for i, title in enumerate(HEADERS, start=2):  # B열부터
        if title:
            ws.cell(8, i, title)
    ws.cell(9, 10, "PO번호")
    ws.cell(9, 11, "SO번호")
    for i, title in enumerate(["OPUN", "CLUN", "OPLO", "CLLO"], start=22):
        ws.cell(9, i, title)
    for col in "BCDEFGHILMNOPQ":
        ws.merge_cells(f"{col}8:{col}9")

    r = 10
    for row in rows:
        ws.cell(r, 1, r - 9)
        c = ws.cell(r, 2, f"{row['no']:04d}")
        c.number_format = row.get("serial_format", f'"N82{yy}"@')
        c = ws.cell(r, 3, row.get("model", "02"))
        c.number_format = '"NTI-"@'
        if row.get("bush"):
            c = ws.cell(r, 4, row["bush"])
            c.number_format = '"Ø"##"x"##"x"##'
        for col, key in ((7, "made"), (8, "inspected"), (16, "due"), (17, "shipped")):
            value = row.get(key)
            if value is None:
                continue
            if date_as_number and isinstance(value, date):
                value = (value - date(1899, 12, 30)).days  # 엑셀 날짜 숫자 (예: 46031)
            ws.cell(r, col, value).number_format = "mm-dd-yy"
        c = ws.cell(r, 9, f"{row['no'] % 1000:03d}")
        c.number_format = f'"NTK{yy}"@'
        c = ws.cell(r, 10, row.get("po", f"{yy}001"))
        c.number_format = '"#"@'
        ws.cell(r, 13, row.get("customer", "가람상수도사업소"))
        ws.cell(r, 14, row.get("final", "가람상수도사업소"))
        ws.cell(r, 15, row.get("project", "가상 사업"))
        ws.cell(r, 18, row.get("note"))
        ws.cell(r, 22, 1900)  # OPUN 측정값
        r += 1
    for _ in range(extra_blank_rows):  # 일련번호 없이 OPUN 숫자만 있는 행 → 제외 대상
        ws.cell(r, 1, r - 9)
        ws.cell(r, 22, 1850)
        r += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
