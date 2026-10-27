from datetime import date

from rest_framework.test import APITestCase

from accounts.models import User
from common.periods import make_period

from .metrics import model_names, monthly_production, production_by_model, production_summary, year_progress
from .models import AnnualProductionTarget, ManufacturingRecord

P = User.Team.PRODUCTION
S = User.Team.SERVICE


def add_records(year, month, count, start=1):
    for i in range(start, start + count):
        ManufacturingRecord.objects.create(
            serial_full=f"N82{str(year)[2:]}{i:04d}",
            serial_short=f"{i:04d}",
            manufactured_date=date(year, month, 10),
        )


class ProductionSummaryTests(APITestCase):
    """specs/09-metrics.md 3절 계산식"""

    def setUp(self):
        AnnualProductionTarget.objects.create(year=2026, target_quantity=120)  # 월 목표 10대
        add_records(2026, 1, 12, start=1)  # 1월 12대
        add_records(2026, 2, 6, start=100)  # 2월 6대
        add_records(2025, 12, 5, start=200)  # 다른 해는 집계에서 빠져야 함

    def test_month_over_target_shortfall_is_zero(self):
        s = production_summary(make_period("month", 2026, month=1))
        self.assertEqual(s["actual"], 12)
        self.assertEqual(s["period_target"], 10.0)
        self.assertEqual(s["achievement_rate"], 120.0)
        self.assertEqual(s["shortfall"], 0)  # 목표 초과 → 부족치 0

    def test_quarter(self):
        s = production_summary(make_period("quarter", 2026, quarter=1))
        self.assertEqual(s["actual"], 18)
        self.assertEqual(s["period_target"], 30.0)  # 120 ÷ 12 × 3
        self.assertEqual(s["achievement_rate"], 60.0)
        self.assertEqual(s["shortfall"], 12.0)

    def test_year(self):
        s = production_summary(make_period("year", 2026))
        self.assertEqual(s["actual"], 18)
        self.assertEqual(s["shortfall"], 102.0)

    def test_no_target_returns_none(self):
        s = production_summary(make_period("year", 2025))
        self.assertEqual(s["actual"], 5)
        self.assertIsNone(s["achievement_rate"])
        self.assertIsNone(s["shortfall"])

    def test_zero_target_has_no_rate(self):
        AnnualProductionTarget.objects.filter(year=2026).update(target_quantity=0)
        s = production_summary(make_period("year", 2026))
        self.assertIsNone(s["achievement_rate"])
        self.assertEqual(s["shortfall"], 0)


class YearOverviewTests(APITestCase):
    """월별 생산량·연간 목표까지 남은 양·연말 예상·모델별 생산량"""

    def setUp(self):
        AnnualProductionTarget.objects.create(year=2026, target_quantity=120)  # 월 목표 10대
        add_records(2026, 1, 12, start=1)
        add_records(2026, 2, 6, start=100)
        add_records(2026, 11, 3, start=200)  # 오늘(10월) 이후 날짜는 누적에서 빠져야 함

    def test_monthly(self):
        m = monthly_production(2026)
        self.assertEqual(len(m["months"]), 12)
        self.assertEqual(m["months"][0], {"month": 1, "actual": 12})
        self.assertEqual(m["months"][2]["actual"], 0)
        self.assertEqual(m["monthly_target"], 10.0)
        self.assertIsNone(monthly_production(2025)["monthly_target"])

    def test_current_year_progress(self):
        p = year_progress(2026, today=date(2026, 10, 27))
        self.assertEqual(p["cumulative"], 18)
        self.assertEqual(p["remaining"], 102)
        self.assertEqual(p["remaining_months"], 3)  # 10·11·12월
        self.assertEqual(p["required_per_month"], 34.0)
        # 18대 ÷ 300일 × 365일 = 21.9 → 22
        self.assertEqual(p["forecast"], 22)

    def test_past_year_has_no_forecast(self):
        p = year_progress(2026, today=date(2027, 3, 1))
        self.assertEqual(p["cumulative"], 21)
        self.assertEqual(p["remaining_months"], 0)
        self.assertIsNone(p["required_per_month"])
        self.assertIsNone(p["forecast"])

    def test_future_year_and_no_target(self):
        p = year_progress(2027, today=date(2026, 10, 27))
        self.assertEqual((p["cumulative"], p["remaining_months"], p["forecast"]), (0, 12, None))
        self.assertIsNone(p["remaining"])  # 2027년 목표 없음

    def test_by_model(self):
        ManufacturingRecord.objects.filter(serial_short__in=["0001", "0002", "0003"]).update(model_name="NTI-04")
        rows = production_by_model(make_period("quarter", 2026, quarter=1))
        self.assertEqual(rows[0], {"name": "(미입력)", "quantity": 15})
        self.assertEqual(rows[1], {"name": "NTI-04", "quantity": 3})
        self.assertEqual(model_names(), ["NTI-04"])  # 중복 없이, 빈 모델명 제외


class ProductionApiTests(APITestCase):
    def login(self, role, view=(), edit=()):
        user = User.objects.create_user(f"u_{role}_{len(view)}{len(edit)}", password="x", role=role,
                                        view_teams=list(view), edit_teams=list(edit))
        self.client.force_authenticate(user)
        return user

    def test_leader_can_set_target(self):
        self.login(User.Role.LEADER, [P], [P])
        res = self.client.put("/api/production/targets/2026/", {"target_quantity": 300}, format="json")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(self.client.get("/api/production/targets/2026/").data["target_quantity"], 300)

    def test_member_can_view_but_not_set_target(self):
        self.login(User.Role.MEMBER, [P])
        self.assertEqual(self.client.get("/api/production/targets/2026/").status_code, 200)
        res = self.client.put("/api/production/targets/2026/", {"target_quantity": 300}, format="json")
        self.assertEqual(res.status_code, 403)

    def test_service_member_cannot_view(self):
        self.login(User.Role.MEMBER, [S])
        self.assertEqual(self.client.get("/api/production/dashboard/").status_code, 403)

    def test_anonymous_401(self):
        self.assertEqual(self.client.get("/api/production/dashboard/").status_code, 401)

    def test_invalid_target_400(self):
        self.login(User.Role.ADMIN)
        res = self.client.put("/api/production/targets/2026/", {"target_quantity": -1}, format="json")
        self.assertEqual(res.status_code, 400)

    def test_dashboard(self):
        self.login(User.Role.LEADER, [P], [P])
        AnnualProductionTarget.objects.create(year=2026, target_quantity=120)
        add_records(2026, 10, 4)
        res = self.client.get("/api/production/dashboard/?year=2026&period=month&month=10")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["period"]["label"], "2026년 10월")
        self.assertEqual(res.data["production"]["actual"], 4)
        self.assertEqual(res.data["production"]["achievement_rate"], 40.0)


class SerialLookupTests(APITestCase):
    """specs/05-production.md 2절 일련번호 조회"""

    def setUp(self):
        user = User.objects.create_user("viewer", password="x", role=User.Role.MEMBER, view_teams=[P])
        self.client.force_authenticate(user)
        ManufacturingRecord.objects.create(
            serial_full="N82260001", serial_short="0001", model_name="NTI-04", bush_machining="301008",
            rpm="29", power="380/3/60", manufactured_date=date(2026, 1, 9), inspection_no="NTK26001",
            po_no="26006", customer="가람상수도사업소", due_date=date(2026, 1, 15), note="방폭",
        )
        ManufacturingRecord.objects.create(serial_full="M82250001", serial_short="0001",
                                           manufactured_date=date(2025, 3, 1))

    def search(self, q):
        return self.client.get("/api/production/serials/", {"q": q}).data["results"]

    def test_search_by_4_digits_returns_all_matches(self):
        results = self.search("0001")
        self.assertEqual([r["serial_full"] for r in results], ["N82260001", "M82250001"])  # 최근 제조 순
        self.assertEqual(len(self.search("1")), 2)  # 1 → 0001

    def test_search_by_full_number(self):
        self.assertEqual([r["serial_full"] for r in self.search("n82260001")], ["N82260001"])
        self.assertEqual(self.search("N82269999"), [])
        self.assertEqual(self.search(""), [])

    def test_detail_has_11_items(self):
        res = self.client.get("/api/production/serials/N82260001/")
        self.assertEqual(res.status_code, 200)
        items = res.data["items"]
        self.assertEqual([i["no"] for i in items], list(range(1, 12)))
        self.assertEqual(items[0]["value"]["부시 가공"], "Ø30x10x08")
        self.assertEqual(items[4]["value"]["PO번호"], "#26006")
        self.assertEqual(items[9]["value"], "2026-01-15")

    def test_detail_not_found_404(self):
        self.assertEqual(self.client.get("/api/production/serials/N00000000/").status_code, 404)

    def test_service_only_user_cannot_search(self):
        other = User.objects.create_user("svc", password="x", role=User.Role.MEMBER, view_teams=[S])
        self.client.force_authenticate(other)
        self.assertEqual(self.client.get("/api/production/serials/", {"q": "0001"}).status_code, 403)


class DueCalendarTests(APITestCase):
    """specs/05-production.md 3절 납기 캘린더"""

    def setUp(self):
        user = User.objects.create_user("viewer", password="x", role=User.Role.MEMBER, view_teams=[P])
        self.client.force_authenticate(user)
        for i, (due, po, customer) in enumerate([
            (date(2026, 10, 5), "26006", "가람"), (date(2026, 10, 5), "26006", "가람"),
            (date(2026, 10, 5), "26007", "나래"), (date(2026, 10, 20), "", "다온"),
            (date(2026, 11, 1), "26008", "라온"),  # 다른 달 → 빠짐
        ], start=1):
            ManufacturingRecord.objects.create(serial_full=f"N8226{i:04d}", serial_short=f"{i:04d}",
                                               due_date=due, po_no=po, customer=customer)

    def test_groups_by_date_and_order_number(self):
        res = self.client.get("/api/production/due-calendar/", {"year": 2026, "month": 10})
        self.assertEqual(res.status_code, 200)
        days = {d["date"]: d["orders"] for d in res.data["days"]}
        self.assertEqual(list(days), ["2026-10-05", "2026-10-20"])
        first = days["2026-10-05"]
        self.assertEqual([(o["po_no"], o["quantity"]) for o in first], [("26006", 2), ("26007", 1)])
        self.assertEqual(first[0]["customers"], ["가람"])
        self.assertEqual(days["2026-10-20"][0]["po_no"], "")  # 수주번호 없음도 표시

    def test_invalid_month_400(self):
        self.assertEqual(self.client.get("/api/production/due-calendar/", {"year": 2026, "month": 13}).status_code, 400)

    def test_service_only_user_403(self):
        self.client.force_authenticate(User.objects.create_user("svc", password="x", view_teams=[S]))
        self.assertEqual(self.client.get("/api/production/due-calendar/").status_code, 403)


class DueStatusTests(APITestCase):
    """납기 처리 상태(출고일 기준)와 메인 대시보드 납기 임박 목록"""

    TODAY = date(2026, 10, 27)

    def setUp(self):
        self.client.force_authenticate(User.objects.create_user("viewer", password="x", role=User.Role.MEMBER,
                                                                view_teams=[P]))
        rows = [  # (납기일, 수주번호, 출고일)
            (date(2026, 10, 10), "26001", date(2026, 10, 8)),                       # 완료
            (date(2026, 10, 20), "26002", None),                                    # 납기 지남·미출고
            (date(2026, 10, 27), "26003", date(2026, 10, 25)), (date(2026, 10, 27), "26003", None),  # 오늘 납기, 일부 출고
            (date(2026, 11, 5), "26004", None), (date(2026, 11, 5), "26004", None),  # D-9, 출고 전
            (date(2026, 11, 20), "26005", None),                                    # 14일 넘음 → 임박 목록 제외
            (date(2026, 9, 1), "26000", None),                                      # 30일 넘게 지남 → 제외
        ]
        for i, (due, po, shipped) in enumerate(rows, start=1):
            ManufacturingRecord.objects.create(serial_full=f"N8226{i:04d}", serial_short=f"{i:04d}", due_date=due,
                                               po_no=po, customer="가상", shipped_date=shipped)

    def test_calendar_status(self):
        from .due_calendar import due_calendar

        days = {d["date"]: d["orders"][0] for d in due_calendar(2026, 10, today=self.TODAY)["days"]}
        self.assertEqual((days["2026-10-10"]["status"], days["2026-10-10"]["overdue"]), ("done", False))
        self.assertEqual((days["2026-10-20"]["status"], days["2026-10-20"]["overdue"]), ("pending", True))
        self.assertEqual((days["2026-10-27"]["status"], days["2026-10-27"]["shipped"], days["2026-10-27"]["overdue"]),
                         ("partial", 1, False))  # 오늘 납기는 아직 지나지 않음

    def test_due_soon(self):
        from .due_calendar import due_soon

        items = due_soon(today=self.TODAY)["items"]
        self.assertEqual([(i["po_no"], i["d_day"], i["status"]) for i in items],
                         [("26002", -7, "pending"), ("26003", 0, "partial"), ("26004", 9, "pending")])
        self.assertEqual(items[2]["quantity"], 2)

    def test_due_soon_api(self):
        res = self.client.get("/api/production/due-soon/")
        self.assertEqual((res.status_code, res.data["days"]), (200, 14))
        self.client.force_authenticate(User.objects.create_user("svc", password="x", view_teams=[S]))
        self.assertEqual(self.client.get("/api/production/due-soon/").status_code, 403)


class DueSearchTests(APITestCase):
    """납기 캘린더 검색: 모든 달에서 수주번호·판매처·사업명·일련번호로 찾기"""

    def setUp(self):
        self.client.force_authenticate(User.objects.create_user("viewer", password="x", role=User.Role.MEMBER,
                                                                view_teams=[P]))
        rows = [  # (일련번호, 납기일, 수주번호, 판매처, 사업명, 출고일)
            ("N82260001", date(2026, 3, 5), "26010", "가람상수도", "가상 정수장 사업", date(2026, 3, 3)),
            ("N82260002", date(2026, 3, 5), "26010", "가람상수도", "가상 정수장 사업", None),
            ("N82260003", date(2026, 11, 2), "26020", "나래발전", "가상 발전 사업", None),
            ("N82250004", date(2025, 7, 1), "25030", "가람상수도", "가상 배관 사업", date(2025, 6, 28)),
        ]
        for serial, due, po, customer, project, shipped in rows:
            ManufacturingRecord.objects.create(serial_full=serial, serial_short=serial[-4:], due_date=due, po_no=po,
                                               customer=customer, project_name=project, shipped_date=shipped)

    def search(self, q):
        return self.client.get("/api/production/due-calendar/search/", {"q": q})

    def test_search_by_customer_across_months_recent_first(self):
        res = self.search("가람")
        self.assertEqual(res.status_code, 200)
        self.assertEqual([(r["due_date"], r["po_no"]) for r in res.data["results"]],
                         [(date(2026, 3, 5), "26010"), (date(2025, 7, 1), "25030")])

    def test_search_by_po_with_hash_and_serial_shows_whole_order(self):
        self.assertEqual([r["po_no"] for r in self.search("#26020").data["results"]], ["26020"])
        order = self.search("N82260002").data["results"][0]  # 일련번호 하나로 찾아도 수주 묶음 전체
        self.assertEqual((order["po_no"], order["quantity"], order["shipped"], order["status"]), ("26010", 2, 1, "partial"))

    def test_search_by_project_and_empty_query(self):
        self.assertEqual([r["po_no"] for r in self.search("발전 사업").data["results"]], ["26020"])
        self.assertEqual(self.search("없는검색어").data["results"], [])
        self.assertEqual(self.search(" # ").status_code, 400)


class SalesSummaryTests(APITestCase):
    """specs/09-metrics.md 3절 판매량 (기준 날짜: 수주서 납기일)"""

    def setUp(self):
        from .models import SalesOrder, SalesOrderItem

        def order(no, industry, customer, due, items):
            o = SalesOrder.objects.create(sales_no=no, industry=industry, customer=customer, due_date=due,
                                          order_doc_no=no[-5:])
            for name, qty, flange in items:
                SalesOrderItem.objects.create(order=o, product_name=name, quantity=qty, flange=flange)

        order("NTK-SO-26006", "상하수도", "가람", date(2026, 10, 3), [("제수 150A", 3, "F10A"), ("제수 200A", 1, "")])
        order("NTK-SO-26007", "발전소", "다온", date(2026, 10, 9), [("제수 150A", 2, "F12")])
        order("NTK-SO-26008", "상하수도", "나래", date(2026, 11, 1), [("제수 150A", 9, "")])  # 11월 → 10월 집계 제외

    def test_rankings(self):
        from .metrics import sales_summary

        s = sales_summary(make_period("month", 2026, month=10))
        self.assertEqual(s["total"], 6)
        pick = lambda rows: [(r["name"], r["quantity"]) for r in rows]
        self.assertEqual(pick(s["by_industry"]), [("상하수도", 4), ("발전소", 2)])
        self.assertEqual(pick(s["by_customer"])[0], ("가람", 4))
        self.assertEqual(pick(s["by_product"]), [("제수 150A", 5), ("제수 200A", 1)])
        self.assertIsNone(s["by_product"][0]["change"])  # 비교 없음

    def test_quarter_includes_november(self):
        from .metrics import sales_summary

        self.assertEqual(sales_summary(make_period("quarter", 2026, quarter=4))["total"], 15)

    def test_flange_supplement_by_order_doc_no(self):
        from .serials import flanges_for

        record = ManufacturingRecord(serial_full="N82260001", po_no="26006")
        self.assertEqual(flanges_for(record), "F10A")
        self.assertIsNone(flanges_for(ManufacturingRecord(serial_full="N82260002", po_no="99999")))


class SalesOptionsTests(APITestCase):
    """판매 실적 설정: 기준 날짜·거래처 기준·제품 묶음·비교 기간·교차 필터"""

    def setUp(self):
        from .models import SalesOrder, SalesOrderItem

        def order(no, industry, customer, due, ordered, items, final=""):
            o = SalesOrder.objects.create(sales_no=no, industry=industry, customer=customer, due_date=due,
                                          order_date=ordered, final_orderer=final)
            for name, qty in items:
                SalesOrderItem.objects.create(order=o, product_name=name, quantity=qty)

        order("A", "상하수도", "가람", date(2026, 10, 3), date(2026, 9, 1),
              [("제수 150A: NTI-03", 3), ("제수 200A: NTI-04", 1)], final="가상 발주처")
        order("B", "발전소", "다온", date(2026, 10, 9), date(2026, 10, 1), [("버터플라이 150A: NTI-05", 2), ("특수 품목", 1)])
        order("C", "상하수도", "가람", date(2026, 9, 20), None, [("제수 150A: NTI-03", 4)])  # 9월 납기, 수주일 없음
        order("D", "상하수도", "가람", date(2025, 10, 5), date(2025, 9, 1), [("제수 150A: NTI-03", 1)])

    def summary(self, **kw):
        from .metrics import sales_summary

        return sales_summary(make_period("month", 2026, month=10), **kw)

    @staticmethod
    def rows(rows):
        return {r["name"]: r["quantity"] for r in rows}

    def test_product_groups(self):
        self.assertEqual(self.rows(self.summary(product_group_by="model")["by_product"]),
                         {"NTI-03": 3, "NTI-04": 1, "NTI-05": 2, "기타(형식 다름)": 1})
        self.assertEqual(self.rows(self.summary(product_group_by="size")["by_product"]),
                         {"150A": 5, "200A": 1, "기타(형식 다름)": 1})
        self.assertEqual(self.rows(self.summary(product_group_by="type")["by_product"]),
                         {"제수": 4, "버터플라이": 2, "기타(형식 다름)": 1})

    def test_date_basis(self):
        s = self.summary(date_basis="order")  # 수주일 10월 = B만
        self.assertEqual(s["total"], 3)
        self.assertEqual(s["missing_date_orders"], 1)  # 수주일이 빈 C
        self.assertEqual(self.summary()["missing_date_orders"], 0)

    def test_customer_basis(self):
        self.assertEqual(self.rows(self.summary(customer_basis="final_orderer")["by_customer"]),
                         {"가상 발주처": 4, "(미입력)": 3})

    def test_cross_filter_keeps_own_ranking(self):
        s = self.summary(industry="상하수도")
        self.assertEqual(s["total"], 4)
        self.assertEqual(self.rows(s["by_industry"]), {"상하수도": 4, "발전소": 3})  # 자기 순위는 전체
        self.assertEqual(self.rows(s["by_customer"]), {"가람": 4})
        s = self.summary(product_group_by="model", product="NTI-03")
        self.assertEqual((s["total"], self.rows(s["by_customer"])), (3, {"가람": 3}))

    def test_compare(self):
        s = self.summary(compare="previous")  # 9월: C 4대
        self.assertEqual(s["compare"], {"label": "2026년 9월", "total": 4, "change": 3})
        garam = [r for r in s["by_customer"] if r["name"] == "가람"][0]
        self.assertEqual((garam["previous"], garam["change"]), (4, 0))
        s = self.summary(compare="last_year")  # 2025년 10월: D 1대
        self.assertEqual(s["compare"]["label"], "2025년 10월")
        self.assertEqual(s["compare"]["total"], 1)

    def test_previous_period_wraps_year(self):
        from .metrics import previous_period

        self.assertEqual(previous_period(make_period("month", 2026, month=1), "previous").label, "2025년 12월")
        self.assertEqual(previous_period(make_period("quarter", 2026, quarter=1), "previous").label, "2025년 4분기")
        self.assertEqual(previous_period(make_period("year", 2026), "previous").label, "2025년")

    def test_api_validates_choices(self):
        user = User.objects.create_user("v", password="x", role=User.Role.MEMBER, view_teams=[P])
        self.client.force_authenticate(user)
        base = "/api/production/sales/?year=2026&period=month&month=10"
        res = self.client.get(f"{base}&product_group=model&compare=last_year")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["sales"]["compare"]["total"], 1)
        self.assertEqual(self.client.get(f"{base}&date_basis=x").status_code, 400)


class ListApiTests(APITestCase):
    """제조내역 목록·수주 내역 목록·엑셀 내려받기"""

    def setUp(self):
        from .models import SalesOrder, SalesOrderItem

        user = User.objects.create_user("viewer", password="x", role=User.Role.MEMBER, view_teams=[P])
        self.client.force_authenticate(user)
        add_records(2026, 10, 3, start=1)
        add_records(2026, 9, 2, start=50)  # 3분기 → 4분기 목록에서 빠져야 함
        ManufacturingRecord.objects.filter(serial_short="0001").update(
            model_name="NTI-04", customer="가람상사", shipped_date=date(2026, 10, 20)
        )
        o = SalesOrder.objects.create(sales_no="NTK-SO-26006", industry="상하수도", customer="가람",
                                      project_name="가상 정수장", due_date=date(2026, 10, 3))
        SalesOrderItem.objects.create(order=o, product_name="제수 150A", quantity=3)
        SalesOrderItem.objects.create(order=o, product_name="제수 200A", quantity=1)
        SalesOrder.objects.create(sales_no="NTK-SO-26007", industry="발전소", due_date=date(2026, 10, 9))
        SalesOrder.objects.create(sales_no="NTK-SO-26008", industry="발전소", due_date=date(2026, 7, 1))

    Q4 = "year=2026&period=quarter&quarter=4"

    def test_records_period_and_counts(self):
        res = self.client.get(f"/api/production/records/?{self.Q4}")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data["total"], 3)
        self.assertEqual(res.data["counts"], {"total": 3, "shipped": 1, "not_shipped": 2})

    def test_records_filters(self):
        get = lambda extra: self.client.get(f"/api/production/records/?{self.Q4}&{extra}").data
        self.assertEqual([r["serial_full"] for r in get("q=가람")["results"]], ["N82260001"])
        self.assertEqual(get("model=NTI-04")["total"], 1)
        self.assertEqual(get("model=(미입력)")["total"], 2)
        shipped_no = get("shipped=no")
        self.assertEqual(shipped_no["total"], 2)
        self.assertEqual(shipped_no["counts"]["total"], 3)  # 버튼 건수는 출고 상태로 거르기 전 기준
        self.assertEqual(self.client.get(f"/api/production/records/?{self.Q4}&shipped=x").status_code, 400)
        self.assertEqual(self.client.get(f"/api/production/records/?{self.Q4}&page=0").status_code, 400)

    def test_records_export(self):
        from io import BytesIO

        from openpyxl import load_workbook

        res = self.client.get(f"/api/production/records/export/?{self.Q4}&shipped=no")
        self.assertEqual(res.status_code, 200)
        self.assertIn("filename*=UTF-8''", res["Content-Disposition"])
        ws = load_workbook(BytesIO(res.content)).active
        self.assertEqual(ws.cell(1, 1).value, "일련번호")
        self.assertEqual(ws.max_row, 3)  # 제목 + 출고 전 2건
        self.assertEqual(ws.cell(2, 6).value, "2026-10-10")  # 날짜는 YYYY-MM-DD 글자

    def test_orders(self):
        data = self.client.get(f"/api/production/orders/?{self.Q4}").data
        self.assertEqual(data["counts"], {"total": 2, "quantity": 4})
        first = [o for o in data["results"] if o["sales_no"] == "NTK-SO-26006"][0]
        self.assertEqual((first["quantity"], len(first["items"])), (4, 2))
        self.assertEqual(self.client.get(f"/api/production/orders/?{self.Q4}&industry=발전소").data["total"], 1)
        # 품명으로 찾아도 수주서가 한 번만 나온다
        self.assertEqual(self.client.get(f"/api/production/orders/?{self.Q4}&q=제수").data["total"], 1)

    def test_orders_export_one_row_per_item(self):
        from io import BytesIO

        from openpyxl import load_workbook

        res = self.client.get(f"/api/production/orders/export/?{self.Q4}")
        ws = load_workbook(BytesIO(res.content)).active
        self.assertEqual(ws.max_row, 4)  # 제목 + 제품 2행 + 제품 없는 수주 1행

    def test_service_member_cannot_view(self):
        user = User.objects.create_user("svc", password="x", role=User.Role.MEMBER, view_teams=[S])
        self.client.force_authenticate(user)
        self.assertEqual(self.client.get(f"/api/production/records/?{self.Q4}").status_code, 403)
        self.assertEqual(self.client.get(f"/api/production/orders/export/?{self.Q4}").status_code, 403)
