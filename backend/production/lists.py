"""생산팀 목록 화면: 제조내역 목록·수주 내역 목록과 엑셀 내려받기 (specs/05-production.md 7절)

- 제조내역 목록: 제조년월이 기간에 속하는 제조내역 (대시보드 생산량과 같은 기준)
- 수주 내역 목록: 납기일이 기간에 속하는 수주서 (대시보드 판매량과 같은 기준)
- 목록과 엑셀은 같은 걸러내기 함수를 써서 화면에 보이는 것과 내려받는 것이 같다
"""

from io import BytesIO
from urllib.parse import quote

from django.db.models import Q, Sum
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from rest_framework.exceptions import ValidationError

from .models import ManufacturingRecord, SalesOrder, SalesOrderItem

PAGE_SIZE = 100
NO_MODEL = "(미입력)"


def page_number(params):
    raw = params.get("page") or "1"
    if not raw.isdigit() or int(raw) < 1:
        raise ValidationError({"page": "page는 1 이상의 숫자여야 합니다."})
    return int(raw)


def paginate(queryset, page):
    total = queryset.count()
    rows = list(queryset[(page - 1) * PAGE_SIZE : page * PAGE_SIZE])
    return rows, {"page": page, "page_size": PAGE_SIZE, "total": total}


# ── 제조내역 ──────────────────────────────────────────────


def filter_records(period, params):
    """기간 + 검색어·모델 → (출고 상태로 거르기 전 목록, 출고 상태까지 거른 목록)"""
    records = ManufacturingRecord.objects.filter(manufactured_date__range=(period.start, period.end))
    q = (params.get("q") or "").strip().lstrip("#")
    if q:
        records = records.filter(
            Q(serial_full__icontains=q)
            | Q(customer__icontains=q)
            | Q(final_orderer__icontains=q)
            | Q(project_name__icontains=q)
            | Q(po_no__icontains=q)
            | Q(customer_order_no__icontains=q)
        )
    model = params.get("model") or ""
    if model:
        records = records.filter(model_name="" if model == NO_MODEL else model)

    shipped = params.get("shipped") or ""
    if shipped not in ("", "yes", "no"):
        raise ValidationError({"shipped": "shipped는 yes 또는 no여야 합니다."})
    filtered = records
    if shipped:
        filtered = records.filter(shipped_date__isnull=shipped == "no")
    return records, filtered.order_by("-manufactured_date", "-serial_full")


def record_counts(records):
    total = records.count()
    shipped = records.filter(shipped_date__isnull=False).count()
    return {"total": total, "shipped": shipped, "not_shipped": total - shipped}


def record_row(r):
    return {
        "serial_full": r.serial_full,
        "model_name": r.model_name,
        "manufactured_date": r.manufactured_date,
        "inspection_date": r.inspection_date,
        "customer": r.customer,
        "project_name": r.project_name,
        "po_no": r.po_no,
        "due_date": r.due_date,
        "shipped_date": r.shipped_date,
    }


RECORD_COLUMNS = [
    ("일련번호", "serial_full"),
    ("모델명", "model_name"),
    ("부시 가공", "bush_machining"),
    ("rpm", "rpm"),
    ("전원", "power"),
    ("제조년월", "manufactured_date"),
    ("검사일자", "inspection_date"),
    ("검사번호", "inspection_no"),
    ("뉴토크 PO번호", "po_no"),
    ("SO번호", "so_no"),
    ("거래처 주문번호", "customer_order_no"),
    ("거래처", "customer"),
    ("최종발주처", "final_orderer"),
    ("사업명", "project_name"),
    ("납기일", "due_date"),
    ("출고일", "shipped_date"),
    ("비고", "note"),
]


# ── 수주 내역 ─────────────────────────────────────────────


def filter_orders(period, params):
    orders = SalesOrder.objects.filter(due_date__range=(period.start, period.end))
    q = (params.get("q") or "").strip().lstrip("#")
    if q:
        orders = orders.filter(
            Q(sales_no__icontains=q)
            | Q(order_doc_no__icontains=q)
            | Q(customer__icontains=q)
            | Q(final_orderer__icontains=q)
            | Q(project_name__icontains=q)
            | Q(items__product_name__icontains=q)
        ).distinct()
    industry = params.get("industry") or ""
    if industry:
        orders = orders.filter(industry=industry)
    return orders.order_by("-due_date", "sales_no")


def order_counts(orders):
    quantity = SalesOrderItem.objects.filter(order__in=orders.values("id")).aggregate(total=Sum("quantity"))["total"]
    return {"total": orders.count(), "quantity": quantity or 0}


def order_row(o):
    items = [
        {
            "product_name": i.product_name,
            "quantity": i.quantity,
            "power": i.power,
            "flange": i.flange,
            "rpm_ot": i.rpm_ot,
            "stem_bush": i.stem_bush,
        }
        for i in o.items.all()
    ]
    return {
        "sales_no": o.sales_no,
        "order_doc_no": o.order_doc_no,
        "order_date": o.order_date,
        "industry": o.industry,
        "business_type_raw": o.business_type_raw,
        "customer": o.customer,
        "final_orderer": o.final_orderer,
        "project_name": o.project_name,
        "delivery_date": o.delivery_date,
        "due_date": o.due_date,
        "note": o.note,
        "items": items,
        "quantity": sum(i["quantity"] for i in items),
    }


# 수주서 1건에 제품이 여러 개면 엑셀에서는 제품 1개 = 1행 (주문 정보는 행마다 반복)
ORDER_COLUMNS = [
    ("판매번호", "sales_no"),
    ("수주서류 번호", "order_doc_no"),
    ("수주일", "order_date"),
    ("산업 분야", "industry"),
    ("사업구분 원문", "business_type_raw"),
    ("판매처", "customer"),
    ("최종발주처/EPC", "final_orderer"),
    ("사업명", "project_name"),
    ("납기일", "due_date"),
    ("납품일/출장일", "delivery_date"),
    ("품명", "item.product_name"),
    ("수량", "item.quantity"),
    ("전원", "item.power"),
    ("Flange", "item.flange"),
    ("RPM/OT", "item.rpm_ot"),
    ("스템/부시가공", "item.stem_bush"),
]


def order_excel_rows(orders):
    for o in orders.prefetch_related("items"):
        for item in o.items.all() or [None]:
            row = {}
            for _, key in ORDER_COLUMNS:
                if key.startswith("item."):
                    row[key] = getattr(item, key[5:]) if item else ""
                else:
                    row[key] = getattr(o, key)
            yield row


# ── 엑셀 파일 ─────────────────────────────────────────────


def excel_response(file_name, sheet_title, columns, rows):
    """columns: [(제목, 키)], rows: 딕셔너리 또는 모델 객체 목록 → .xlsx 내려받기 응답"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append([title for title, _ in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        values = [row.get(key) if isinstance(row, dict) else getattr(row, key) for _, key in columns]
        # 날짜는 YYYY-MM-DD 글자로 (CLAUDE.md 6절 형식 통일), 빈 값은 빈칸
        ws.append([v.isoformat() if hasattr(v, "isoformat") else ("" if v is None else v) for v in values])
    for i, (title, _) in enumerate(columns, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(10, len(title) * 2 + 2)
    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    # 한글 파일 이름은 filename* (UTF-8) 로 보낸다
    response["Content-Disposition"] = f"attachment; filename=\"export.xlsx\"; filename*=UTF-8''{quote(file_name)}"
    return response
